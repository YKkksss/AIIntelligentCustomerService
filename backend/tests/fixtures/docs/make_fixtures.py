"""生成固定样例文件（已提交到仓库；只在需要修改样例时手动运行）。

    uv run python tests/fixtures/docs/make_fixtures.py

内容为虚构的电商售后政策，不含任何真实数据。重新生成后需同步更新快照：
    AICS_UPDATE_SNAPSHOTS=1 uv run pytest tests/unit/test_parsers_snapshot.py
"""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf

HERE = Path(__file__).parent

POLICY = [
    ("h1", "售后服务政策"),
    ("p", "本政策适用于小优商城自营商品。如与活动页面说明不一致，以活动页面为准。"),
    ("h2", "退货规则"),
    ("p", "签收后 7 天内，商品完好且不影响二次销售的，可申请无理由退货。定制类、生鲜类商品除外。"),
    ("p", "因质量问题退货的，运费由商家承担；非质量问题退货的，退回运费由买家承担。"),
    ("h2", "换货规则"),
    ("p", "签收后 15 天内出现性能故障，可申请换货。换货商品发出后，原商品需在 7 天内寄回。"),
    (
        "table",
        [
            ["商品类别", "退货时限", "换货时限"],
            ["数码家电", "7 天", "15 天"],
            ["服饰鞋包", "7 天", "30 天"],
        ],
    ),
    ("h2", "退款时效"),
    (
        "p",
        "仓库签收退货后 1 个工作日内审核，审核通过后原路退款，到账时间以银行为准，一般为 1–5 个工作日。",
    ),
]


def make_md() -> None:
    out = ["---", "author: demo", "---", ""]
    for kind, v in POLICY:
        if kind == "h1":
            out += [f"# {v}", ""]
        elif kind == "h2":
            out += [f"## {v}", ""]
        elif kind == "p":
            out += [str(v), ""]
        else:
            rows = v  # type: ignore[assignment]
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + " --- |" * len(rows[0]))
            out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            out.append("")
    out += [
        "- 联系客服：在线客服 9:00–21:00",
        "- 客服电话：400-000-0000",
        "",
        "```",
        "# 这不是标题",
        "```",
        "",
    ]
    (HERE / "sample.md").write_text("\n".join(out), encoding="utf-8")


def make_txt() -> None:
    paras = [str(v) for k, v in POLICY if k in ("h1", "h2", "p")]
    (HERE / "sample_gbk.txt").write_bytes("\n\n".join(paras).encode("gbk"))


def make_html() -> None:
    body = []
    for kind, v in POLICY:
        if kind == "h1":
            body.append(f"<h1>{v}</h1>")
        elif kind == "h2":
            body.append(f"<h2>{v}</h2>")
        elif kind == "p":
            body.append(f"<p>{v}</p>")
        else:
            rows = v  # type: ignore[assignment]
            trs = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in rows)
            body.append(f"<table>{trs}</table>")
    html = (
        "<html><head><meta charset='utf-8'><title>售后服务政策 - 小优商城</title></head><body>"
        "<nav><a href='/'>首页</a> | <a href='/help'>帮助中心</a></nav>"
        f"<article>{''.join(body)}</article>"
        "<footer>© 2026 小优商城 版权所有</footer><script>var x=1;</script></body></html>"
    )
    (HERE / "sample.html").write_text(html, encoding="utf-8")


def make_docx() -> None:
    d = docx.Document()
    d.core_properties.title = "售后服务政策"
    for kind, v in POLICY:
        if kind == "h1":
            d.add_heading(str(v), level=1)
        elif kind == "h2":
            d.add_heading(str(v), level=2)
        elif kind == "p":
            d.add_paragraph(str(v))
        else:
            rows = v  # type: ignore[assignment]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
    d.add_paragraph("在线客服 9:00–21:00", style="List Bullet")
    d.add_paragraph("客服电话：400-000-0000", style="List Bullet")
    d.save(HERE / "sample.docx")


def make_pdf() -> None:
    doc = pymupdf.open()
    doc.set_metadata({"title": "售后服务政策"})
    font = "china-s"
    items = POLICY + [("p", "其他未尽事宜，请联系在线客服咨询。")] * 3
    page = None
    y = 0.0
    pno = 0

    def new_page() -> None:
        nonlocal page, y, pno
        page = doc.new_page(width=595, height=842)
        pno += 1
        page.insert_text((60, 30), "小优商城 售后服务政策（内部资料）", fontname=font, fontsize=8)
        page.insert_text((280, 825), f"第 {pno} 页", fontname=font, fontsize=8)
        y = 90.0

    new_page()
    for idx, (kind, v) in enumerate(items):
        if idx in (5, 9):  # 强制分页，让页眉页脚重复出现
            new_page()
        assert page is not None
        if kind in ("h1", "h2"):
            size = 20 if kind == "h1" else 15
            page.insert_text((60, y), str(v), fontname=font, fontsize=size)
            y += size + 14
        elif kind == "p":
            rect = pymupdf.Rect(60, y - 10, 535, y + 60)
            page.insert_textbox(rect, str(v), fontname=font, fontsize=10.5)
            y += 52
        else:
            rows = v  # type: ignore[assignment]
            x0, w, h = 60.0, 150.0, 22.0
            top = y
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    cell = pymupdf.Rect(
                        x0 + j * w, top + i * h, x0 + (j + 1) * w, top + (i + 1) * h
                    )
                    page.draw_rect(cell, color=(0, 0, 0), width=0.6)
                    page.insert_text((cell.x0 + 6, cell.y1 - 7), c, fontname=font, fontsize=10)
            y = top + len(rows) * h + 30
    doc.save(HERE / "sample.pdf", garbage=4, deflate=True, no_new_id=True)


if __name__ == "__main__":
    make_md()
    make_txt()
    make_html()
    make_docx()
    make_pdf()
    print("ok")
